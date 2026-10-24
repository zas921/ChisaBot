"""Bounded, non-executing text extraction for chat attachments. Bytes arrive on stdin."""
import gzip
import bz2
import lzma
import codecs
import io
import json
import os
import re
import shutil
import subprocess
import sys
import tarfile
import tempfile
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

MAX_INPUT = 20 * 1024 * 1024
MAX_TEXT = 120_000
MAX_XML = 8 * 1024 * 1024
MAX_EXPANDED = 64 * 1024 * 1024
MAX_ENTRIES = 10_000


class InvalidFile(Exception):
    pass


class Text:
    def __init__(self):
        self.parts = []
        self.size = 0
        self.truncated = False

    def add(self, value):
        value = str(value).replace("\x00", "")
        if not value:
            return
        remaining = MAX_TEXT - self.size
        if len(value) + 1 > remaining:
            self.truncated = True
        if remaining > 0:
            chunk = (value + "\n")[:remaining]
            self.parts.append(chunk)
            self.size += len(chunk)

    def result(self):
        text = "".join(self.parts).strip()
        if self.truncated:
            text += "\n\n[附件内容超过 120000 字符，以上为部分内容，不能据此断言全文统计结果。]"
        return text


def xml(data):
    if len(data) > MAX_XML or re.search(br"<!\s*(?:DOCTYPE|ENTITY)\b", data, re.I):
        raise InvalidFile("文档 XML 过大或包含不支持的声明。")
    return ET.fromstring(data)


def checked_zip(data):
    archive = zipfile.ZipFile(io.BytesIO(data))
    entries = archive.infolist()
    if len(entries) > MAX_ENTRIES or sum(i.file_size for i in entries) > MAX_EXPANDED:
        raise InvalidFile("压缩后的文档展开体积过大，请拆分文件。")
    if any(i.flag_bits & 1 for i in entries):
        raise InvalidFile("文件已加密，请先去除密码再上传。")
    return archive


def read_xml(archive, name):
    if archive.getinfo(name).file_size > MAX_XML:
        raise InvalidFile("文档的单个 XML 部件过大。")
    with archive.open(name) as stream:
        return xml(stream.read(MAX_XML + 1))


def words(element):
    result = []
    for item in element.iter():
        tag = item.tag.rsplit("}", 1)[-1]
        if tag in ("t", "text") and item.text:
            result.append(item.text)
        elif tag == "tab":
            result.append("\t")
        elif tag in ("br", "cr"):
            result.append("\n")
    return "".join(result)


def external(command, data, extension):
    executable = shutil.which(command)
    if not executable:
        raise InvalidFile("服务器缺少此格式的读取组件，请联系管理员更新服务器。")
    with tempfile.TemporaryDirectory(prefix="chisabot-attachment-") as directory:
        source = Path(directory) / ("input" + extension)
        source.write_bytes(data)
        args = [executable]
        if command == "antiword":
            args += ["-m", "UTF-8.txt", str(source)]
        elif command in ("catdoc", "xls2csv", "catppt"):
            args += ["-d", "utf-8", str(source)]
        else:
            args += ["l", "-slt", "--", str(source)]
        with tempfile.TemporaryFile() as output, tempfile.TemporaryFile() as errors:
            run = subprocess.run(args, stdin=subprocess.DEVNULL, stdout=output, stderr=errors,
                                 timeout=12, check=False, env={**os.environ, "LANG": "C.UTF-8"})
            if run.returncode != 0:
                raise InvalidFile("文档无法读取，可能已损坏、加密或不是所声明的格式。")
            output.seek(0)
            return output.read(MAX_TEXT * 4 + 1).decode("utf-8", errors="replace").replace(str(source), "附件")


def extract(data, extension, mime):
    out = Text()
    if data.startswith(b"%PDF-"):
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise InvalidFile("PDF 已加密，请先去除密码。")
        for number, page in enumerate(reader.pages):
            if number >= 500 or out.size >= MAX_TEXT:
                out.truncated = True
                break
            out.add(f"[第 {number + 1} 页]\n{page.extract_text() or ''}")
        if not any(part.strip() for part in re.sub(r"\[第 \d+ 页\]", "", out.result()).splitlines()):
            out.add("此 PDF 未提取到文字，可能是扫描件。尚未进行 OCR，不能推断页面内容；请提供清晰图片或带文字层的 PDF。")
    elif data.startswith(b"PK"):
        with checked_zip(data) as archive:
            names = archive.namelist()
            if "word/document.xml" in names:
                document = read_xml(archive, "word/document.xml")
                body = next((e for e in document if e.tag.endswith("}body")), document)
                for element in body:
                    if out.size >= MAX_TEXT:
                        out.truncated = True
                        break
                    if element.tag.endswith("}tbl"):
                        for row in element:
                            out.add("\t".join(words(cell) for cell in row))
                    else:
                        out.add(words(element))
                for name in names:
                    if re.fullmatch(r"word/(?:header\d+|footer\d+|footnotes|endnotes)\.xml", name) and out.size < MAX_TEXT:
                        for p in read_xml(archive, name).iter():
                            if p.tag.endswith("}p"):
                                out.add(words(p))
            elif "xl/workbook.xml" in names:
                from openpyxl import load_workbook
                book = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
                try:
                    for sheet in book:
                        out.add(f"[工作表：{sheet.title}；公式保留原文，未重新计算]")
                        if sheet.max_row and sheet.max_row > 5000 or sheet.max_column and sheet.max_column > 100:
                            out.truncated = True
                        for row_number, row in enumerate(sheet.iter_rows(max_row=min(sheet.max_row or 5000, 5000), max_col=min(sheet.max_column or 100, 100)), 1):
                            if out.size >= MAX_TEXT:
                                out.truncated = True
                                break
                            values = [f"{cell.column_letter}={cell.value}" for cell in row if cell.value is not None]
                            if values:
                                out.add(f"{row_number}: " + "\t".join(values))
                        if out.size >= MAX_TEXT:
                            break
                finally:
                    book.close()
            elif any(re.fullmatch(r"ppt/slides/slide\d+\.xml", name) for name in names):
                slides = sorted((n for n in names if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)), key=lambda n: int(re.search(r"(\d+)\.xml$", n)[1]))
                for index, name in enumerate(slides):
                    if out.size >= MAX_TEXT:
                        out.truncated = True
                        break
                    out.add(f"[幻灯片 {index + 1}]")
                    for element in read_xml(archive, name).iter():
                        if element.tag.endswith("}p"):
                            out.add(words(element))
                    note = name.replace("slides/slide", "notesSlides/notesSlide")
                    if note in names:
                        for element in read_xml(archive, note).iter():
                            if element.tag.endswith("}p"):
                                out.add(words(element))
            elif "content.xml" in names and "mimetype" in names:
                for element in read_xml(archive, "content.xml").iter():
                    if element.tag.rsplit("}", 1)[-1] in ("p", "h"):
                        out.add("".join(element.itertext()))
                    if out.size >= MAX_TEXT:
                        out.truncated = True
                        break
            else:
                if extension in (".docx", ".xlsx", ".pptx", ".odt", ".ods", ".odp"):
                    raise InvalidFile("文件中缺少对应的办公文档内容，请检查文件格式。")
                out.add("[压缩包目录；仅列出文件，尚未读取内部文件正文，未执行或释放其中内容]")
                for item in archive.infolist()[:200]:
                    out.add(f"{item.filename} ({item.file_size} bytes)")
                if len(names) > 200:
                    out.truncated = True
    elif data.startswith(bytes.fromhex("d0cf11e0a1b11ae1")):
        if extension in (".docx", ".xlsx", ".pptx", ".docm", ".xlsm", ".pptm"):
            raise InvalidFile("Office 文件已加密，请先去除密码再上传。")
        command = {".doc": "antiword", ".xls": "xls2csv", ".ppt": "catppt"}.get(extension)
        if not command:
            raise InvalidFile("旧版 Office 文件需要正确的 .doc、.xls 或 .ppt 扩展名。")
        out.add(external(command, data, extension))
    elif data.startswith(b"{\\rtf"):
        out.add(external("catdoc", data, ".rtf"))
    elif extension in (".7z", ".rar"):
        out.add("[压缩包目录；尚未读取内部文件正文]")
        out.add(external("7z", data, extension))
    elif extension in (".tar", ".tgz", ".gz", ".bz2", ".xz"):
        try:
            with tarfile.open(fileobj=io.BytesIO(data), mode="r|*") as archive:
                out.add("[压缩包目录；尚未读取内部文件正文]")
                for index, entry in enumerate(archive):
                    if index >= 200:
                        out.truncated = True
                        break
                    out.add(f"{entry.name} ({entry.size} bytes)")
        except tarfile.ReadError:
            decompressor = {".gz": gzip.GzipFile, ".bz2": bz2.BZ2File, ".xz": lzma.LZMAFile}.get(extension)
            if decompressor is None:
                raise InvalidFile("压缩文件无法读取。")
            compressed = gzip.GzipFile(fileobj=io.BytesIO(data)) if extension == ".gz" else decompressor(io.BytesIO(data))
            with compressed as stream:
                content = stream.read(MAX_TEXT * 4 + 1)
            out.add(codecs.getincrementaldecoder("utf-8")("strict").decode(content, final=len(content) <= MAX_TEXT * 4))
    else:
        out.add(f"此附件的类型为 {mime}，原文件可以下载。未提取其音视频或二进制内容，不能声称已听取、观看或读取正文。")
    result = out.result()
    if not result:
        raise InvalidFile("未提取到可读文字，文件可能为空、已加密或只包含图片。")
    return {"text": result, "truncated": out.truncated}


def main():
    try:
        import resource
        resource.setrlimit(resource.RLIMIT_CPU, (15, 15))
        resource.setrlimit(resource.RLIMIT_AS, (512 * 1024 * 1024, 512 * 1024 * 1024))
        resource.setrlimit(resource.RLIMIT_FSIZE, (8 * 1024 * 1024, 8 * 1024 * 1024))
    except (ImportError, ValueError, OSError):
        pass
    data = sys.stdin.buffer.read(MAX_INPUT + 1)
    try:
        if len(data) > MAX_INPUT:
            raise InvalidFile("附件超过 20 MB，请拆分后上传。")
        result = extract(data, sys.argv[1].lower(), sys.argv[2])
    except InvalidFile as error:
        result = {"error": str(error)}
    except (zipfile.BadZipFile, ET.ParseError, KeyError, UnicodeError, ValueError):
        result = {"error": "文件无法解析，可能已损坏、加密或格式不正确。"}
    except Exception:
        result = {"error": "文件读取失败，请检查文件是否损坏、加密或过于复杂。"}
    sys.stdout.buffer.write(json.dumps(result, ensure_ascii=False).encode("utf-8"))


if __name__ == "__main__":
    main()
